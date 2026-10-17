"""paperpull_core.sheets, the rule that keeps text a spreadsheet tool writes
from becoming a formula, and the census that every tool writing a
spreadsheet for a person goes through it. The write paths of the two tools
are tested in test_export_purchases.py and test_export_transactions.py.

The same rule for the apps' own CSV files, which are read back as well, so
every value has to read back as it was written, and the census that every
reader and writer of a CSV in the repository goes through it. CsvFile is
tested in test_csv.py and every app's own write, read and Rename in
test_every_app_keeps_its_csv_text_as_text.py. Every value here is made up."""
import ast
import csv
import functools
import importlib.util
import io
import itertools
import os
import random
import re
import subprocess
import sys
import unicodedata
from pathlib import Path

import pytest

from paperpull_core.sheets import (csv_value, keep_text, read_stored, read_stored_row,
                                   starts_a_formula, store_value, whole_characters, xlsx_value)

TOOLS = Path(__file__).resolve().parents[2] / "tools"

FORMULAS = [
    "=1+1", "+1+1", "-1+1", "@SUM(1,1)",
    '=HYPERLINK("https://example.invalid/?"&A1,"open")',
    "\t=1+1", "\r=1+1", "\n=1+1", "\tplain words after a tab",
    " =1+1", "   @SUM(1,1)", " =1+1",
    "-", "+", "=", "@", "-.5", "-1e3", "-1,2", "-12.34x", "--12", "+12.34",
    # A money sign that is letters can start a cell name, -kr12 is cell KR12,
    # and Excel opened -kr12.34 and -zł12.34 as formulas. So did -$A$1.
    "-$A$1", "-$B2", "-kr12.34", "-zł12.34", "-$12.34x", "-$ 12.34", "-$.5", "-$-1",
    # Full-width and small signs, which OWASP names for Japanese settings,
    # and a minus before digits that are not ASCII ones.
    "＝1+1", "＋1+1", "－1+1", "＠SUM(1,1)", "﹦1+1", "﹢1+1",
    "－12", "-١٢٣",
]
TEXT = ["", "Kettle", "a=b", "1-1", "12.34", "1,234.56", "'=1+1", "(12.34)",
        "PAYMENT - THANK YOU", "AplPay EXAMPLE STORE", "Size M; 20% off", "a;b",
        "ＡＢＣ",
        # A whole value starting with " is quoted by the csv module, and both
        # a comma reader and one splitting at ; read it as the text it is.
        '"=1+1']
PLAIN_NEGATIVE = ["-12.34", "-1,234.56", "-42", "-0.5", "-1,000,000", " -12.34",
                  # A refund as the apps write one. Excel opened each as a
                  # number, -£12.34 as text, none as a formula.
                  "-$12.34", "-$1,234.56", "-$0.99", "-$12", "-€12.34", "-£12.34"]
# A value with a ; or a line break inside, and the CSV cell it becomes.
PIECES = [
    ("Widget;=1+1", "Widget;'=1+1"),
    ("Widget; @SUM(1,1)", "Widget;' @SUM(1,1)"),
    ("Widget\n=1+1", "Widget\n'=1+1"),
    ("Widget\r\n+1+1", "Widget\r\n'+1+1"),
    ("Widget\r-1+1", "Widget\r'-1+1"),
    ("=1+1;=2+2", "'=1+1;'=2+2"),
    ("\n=1+1", "'\n'=1+1"),
    ("\r=1+1", "'\r'=1+1"),
    ("Refund; -12.34", "Refund;' -12.34"),
    ("q;-1", "q;'-1"),
    ("Size M; 20% off; final sale", "Size M; 20% off; final sale"),
    ("a;;b;", "a;;b;"),
    # A piece starting with " is held as "" inside the quoted value, which a
    # reader splitting at ; takes for an empty quoted part before the rest.
    ('x;"=1+1;y', 'x;\'"=1+1;y'),
    ('x\n"@SUM(1,1)', 'x\n\'"@SUM(1,1)'),
    ('x; "-1+1', 'x;\' "-1+1'),
    ("Widget;'=1+1", "Widget;''=1+1"),
    ('12" deluxe; "special" edition', '12" deluxe; "special" edition'),
]
# Excel drops a NUL from a CSV it opens, and =1+1 behind one opened as a
# working formula. openpyxl takes no NUL in a cell, so these are CSV only.
BEHIND_A_NUL = [("\x00=1+1", "'\x00=1+1"), ("\x00@SUM(1,1)", "'\x00@SUM(1,1)"),
                ("x;\x00-1+1", "x;'\x00-1+1"), ("\x00-12.34", "\x00-12.34")]


@pytest.mark.parametrize("text, cell", BEHIND_A_NUL)
def test_a_nul_in_front_hides_nothing(text, cell):
    assert csv_value(text) == cell
    assert starts_a_formula(text) is cell.startswith("'")


@pytest.mark.parametrize("text", FORMULAS)
def test_text_that_could_start_a_formula_is_caught(text):
    assert starts_a_formula(text)
    if any(b in text for b in ";\r\n"):
        assert csv_value(text).startswith("'")       # and more, in PIECES
    else:
        assert csv_value(text) == "'" + text


@pytest.mark.parametrize("text", TEXT + PLAIN_NEGATIVE)
def test_other_text_and_a_plain_negative_number_are_left_alone(text):
    assert not starts_a_formula(text)
    assert csv_value(text) == text


@pytest.mark.parametrize("text, cell", PIECES)
def test_a_piece_after_a_semicolon_or_a_line_break_gets_its_own_quote(text, cell):
    """Excel in a country whose list separator is ; starts a new cell at
    each ; and a new row at a line break inside a quoted value."""
    assert csv_value(text) == cell


def test_a_number_is_written_as_a_number_and_none_as_an_empty_cell():
    assert csv_value(-12.34) == -12.34
    assert csv_value(3) == 3
    assert csv_value(0.0) == 0.0
    assert csv_value(None) == ""


def test_an_xlsx_cell_holding_text_is_a_string_cell_with_its_value_unchanged():
    openpyxl = pytest.importorskip("openpyxl")
    ws = openpyxl.Workbook().active
    cells = {}
    for i, value in enumerate(FORMULAS + TEXT + PLAIN_NEGATIVE + ["#N/A", "#DIV/0!"], start=1):
        cell = ws.cell(row=i, column=1)
        cell.value = value
        keep_text(cell)
        cells[value] = cell
    for value, cell in cells.items():
        assert cell.data_type == "s" and cell.value == value, value
        assert cell.quotePrefix is starts_a_formula(value), value


def test_an_xlsx_number_and_an_empty_cell_are_left_as_they_were():
    openpyxl = pytest.importorskip("openpyxl")
    ws = openpyxl.Workbook().active
    amount, count, empty = ws["A1"], ws["A2"], ws["A3"]
    amount.value, count.value = -12.34, 7
    for cell in (amount, count, empty):
        keep_text(cell)
    assert (amount.data_type, amount.value, amount.quotePrefix) == ("n", -12.34, False)
    assert (count.data_type, count.value) == ("n", 7)
    assert empty.value is None and empty.data_type == "n"


# -- what an xlsx cell cannot hold ---------------------------------------------------

def _xml_can_hold(n: int) -> bool:
    """XML 1.0's Char production, written out from the standard, so the
    tests below are not the rule checking itself."""
    return (n in (0x9, 0xA, 0xD) or 0x20 <= n <= 0xD7FF or 0xE000 <= n <= 0xFFFD
            or 0x10000 <= n <= 0x10FFFF)


# Every character an xlsx cannot hold, the 29 control characters openpyxl
# refuses, U+FFFE, U+FFFF and the 2,048 surrogates.
CANNOT_HOLD = [chr(n) for n in range(0x110000) if not _xml_can_hold(n)]
# The six of them Python counts as blank space, which stand between words.
BLANK = [c for c in CANNOT_HOLD if c.isspace()]
# The surrogates, each half of a character UTF-16 writes as two, and the
# mark the browser shows for one on its own (HALF OF A CHARACTER in sheets).
HALVES = frozenset(chr(n) for n in range(0xD800, 0xE000))
MARK = chr(0xFFFD)
HIGH, LOW = chr(0xD83D), chr(0xDE00)          # the two halves of U+1F600


def _made_whole(text: str) -> str:
    """The text read as UTF-16 reads it, written out from the Unicode
    Standard's definition of a surrogate pair, not from Python's codec. A
    high half followed by a low one is the character the two make, and a
    half on its own is U+FFFD."""
    out, i = [], 0
    while i < len(text):
        n = ord(text[i])
        after = ord(text[i + 1]) if i + 1 < len(text) else 0
        if 0xD800 <= n <= 0xDBFF and 0xDC00 <= after <= 0xDFFF:
            out.append(chr(0x10000 + ((n - 0xD800) << 10) + (after - 0xDC00)))
            i += 2
            continue
        out.append(MARK if 0xD800 <= n <= 0xDFFF else text[i])
        i += 1
    return "".join(out)


def _held(text: str) -> str:
    """The text as a cell holds it, made whole, then each character an xlsx
    cannot hold left out, a space for one that is blank space."""
    return "".join(c if _xml_can_hold(ord(c)) else " " if c.isspace() else ""
                   for c in _made_whole(text))


def _excel_reads(text: str) -> str:
    """What Excel shows for an xlsx string. Measured 2026-10-16, _x and four
    hex digits of either case and _ are one character, read once from the
    left, so _x005F_x003D_ shows as _x003D_. _X003D_, _x003D and _x03D_
    show as they are."""
    return re.sub(r"_x([0-9A-Fa-f]{4})_", lambda m: chr(int(m.group(1), 16)), text)


def test_what_an_xlsx_cannot_hold_takes_in_all_openpyxl_refuses():
    pytest.importorskip("openpyxl")
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    refused = [c for c in CANNOT_HOLD if ILLEGAL_CHARACTERS_RE.search(c)]
    assert len(refused) == 29 and len(CANNOT_HOLD) == 29 + 2 + 2048
    assert not ILLEGAL_CHARACTERS_RE.search(_held("".join(map(chr, range(0x110000)))))


def test_every_character_an_xlsx_cannot_hold_is_dropped_a_space_or_the_mark_and_no_other():
    assert BLANK == [chr(0x0B), chr(0x0C), chr(0x1C), chr(0x1D), chr(0x1E), chr(0x1F)]
    every = "".join(map(chr, range(0x110000)))
    assert xlsx_value(every) == _held(every)
    for c in CANNOT_HOLD:
        between = " " if c in BLANK else MARK if c in HALVES else ""
        assert xlsx_value("Kettle" + c + "Lid") == "Kettle" + between + "Lid", ascii(c)
        assert xlsx_value(c * 3) == between * 3
    assert xlsx_value(-12.34) == -12.34 and xlsx_value(7) == 7 and xlsx_value(None) is None
    for text in FORMULAS + TEXT + PLAIN_NEGATIVE:
        assert xlsx_value(text) == text, text


def test_a_character_taken_out_from_in_front_hides_no_formula_and_the_workbook_opens(tmp_path):
    """openpyxl refused each of these as the row went in, and wrote U+FFFE,
    U+FFFF and a surrogate into a workbook neither it nor Excel would open.
    Taken out, each leaves the text around it, and what is left gets the
    quote prefix when it starts like a formula. A half becomes U+FFFD,
    which Excel showed in front of =1+1 and @SUM(1,1) as text, on opening
    and after an edit (measured 2026-10-16), so that needs no prefix."""
    openpyxl = pytest.importorskip("openpyxl")
    wb = openpyxl.Workbook()
    ws = wb.active
    for c in CANNOT_HOLD:
        ws.append([xlsx_value(v) for v in (c + "=1+1", "=" + c + "1+1", " " + c + "@SUM(1,1)",
                                          "Kettle" + c + "lid", c + "-12.34")])
    for row in ws.iter_rows():
        for cell in row:
            keep_text(cell)
    wb.save(tmp_path / "t.xlsx")
    back = openpyxl.load_workbook(tmp_path / "t.xlsx").active
    rows = [[(c.value, c.data_type, c.quotePrefix) for c in row] for row in back.iter_rows()]
    assert len(rows) == len(CANNOT_HOLD)
    for c, row in zip(CANNOT_HOLD, rows):
        s = " " if c in BLANK else MARK if c in HALVES else ""
        front = s != MARK
        assert row == [(s + "=1+1", "s", front), ("=" + s + "1+1", "s", True),
                       (" " + s + "@SUM(1,1)", "s", front), ("Kettle" + s + "lid", "s", False),
                       (s + "-12.34", "s", False)], ascii(c)


ESCAPED = [
    ("_x003D_1+1", "_x005F_x003D_1+1"),
    ("_x003d_1+1", "_x005F_x003d_1+1"),
    ("a_x0041_b", "a_x005F_x0041_b"),
    ("__x003D_1+1", "__x005F_x003D_1+1"),
    ("_x005F_x003D_1+1", "_x005F_x005F_x005F_x003D_1+1"),
    ("_x0041__x0042_", "_x005F_x0041__x005F_x0042_"),
    ("_x0000_", "_x005F_x0000_"),
    # Excel showed these as they are.
    ("_X003D_1+1", "_X003D_1+1"), ("_x003D1+1", "_x003D1+1"), ("_x03D_1+1", "_x03D_1+1"),
    ("_x0003D_1+1", "_x0003D_1+1"), ("_xG03D_1+1", "_xG03D_1+1"),
    # A character dropped from inside the run comes out before the run is seen.
    ("_x" + chr(0) + "003D_1+1", "_x005F_x003D_1+1"),
    ("_" + chr(0xFFFF) + "x003D_1+1", "_x005F_x003D_1+1"),
    # A half there is the mark, which joins nothing up.
    ("_x" + chr(0xD800) + "003D_1+1", "_x" + chr(0xFFFD) + "003D_1+1"),
]


@pytest.mark.parametrize("text, cell", ESCAPED)
def test_text_excel_would_read_as_its_escape_shows_as_it_came(text, cell):
    """Excel read _x003D_1+1 in an xlsx as =1+1, which keep_text had not
    given the quote prefix, and an edit made it a working formula."""
    assert xlsx_value(text) == cell
    assert _excel_reads(cell) == _held(text)


def test_excel_shows_every_text_as_it_came_and_quoted_where_what_it_shows_is_a_formula():
    """A model of the reader, not the rule again. Over made-up texts built
    from the pieces that matter, Excel shows the cell as the text less what
    an xlsx cannot hold, and the cell has the quote prefix exactly when what
    Excel shows starts like a formula."""
    openpyxl = pytest.importorskip("openpyxl")
    pieces = ["_", "x", "X", "0", "3", "D", "d", "5", "F", "=", "1", "+", "@", " ", "a",
              "\x00", "\x01", "\x0b", "\x1f", chr(0xFFFF), chr(0xD800), HIGH, LOW, MARK,
              "_x003D_", "_x005F_", "_x00", "_x0040_"]
    rng = random.Random(20261016)
    ws = openpyxl.Workbook().active
    for i in range(1, 20001):
        text = "".join(rng.choice(pieces) for _ in range(rng.randint(0, 8)))
        cell = ws.cell(row=i, column=1)
        cell.value = xlsx_value(text)
        keep_text(cell)
        shown = _excel_reads(cell.value)
        assert shown == _held(text), ascii(text)
        assert cell.quotePrefix is starts_a_formula(shown), ascii(text)


# -- half of a character -------------------------------------------------------------

def test_a_half_on_its_own_is_the_mark_and_two_side_by_side_are_one_character():
    """Chromium showed a lone half as U+FFFD and printed it into the PDF as
    U+FFFD, and showed a high half then a low one as the character the two
    make (measured 2026-10-16)."""
    for c in HALVES:
        assert whole_characters("SHOP " + c + "NAME") == "SHOP " + MARK + "NAME", ascii(c)
        assert whole_characters(c * 3) == MARK * 3, ascii(c)
    for high in (0xD800, 0xD83D, 0xDBFF):
        for low in (0xDC00, 0xDE00, 0xDFFF):
            pair = chr(high) + chr(low)
            assert whole_characters(pair) == _made_whole(pair) and len(_made_whole(pair)) == 1
    assert whole_characters(HIGH + LOW) == chr(0x1F600)
    assert whole_characters(LOW + HIGH) == MARK * 2
    assert whole_characters(HIGH + HIGH + LOW) == MARK + chr(0x1F600)
    assert whole_characters("Great gift " + chr(0xD83C)) == "Great gift " + MARK


def test_every_text_is_made_whole_as_utf16_reads_it():
    """Against the definition written out in _made_whole, over every
    character in one text and 20,000 made-up texts of halves and what
    stands around them. What comes out always goes into UTF-8."""
    every = "".join(map(chr, range(0x110000)))
    assert whole_characters(every) == _made_whole(every)
    pieces = [HIGH, LOW, chr(0xD800), chr(0xDBFF), chr(0xDC00), chr(0xDFFF), "a", "=", " ",
              ";", "'", MARK, chr(0x1F600), "\x00"]
    rng = random.Random(20261016)
    for _ in range(20000):
        text = "".join(rng.choice(pieces) for _ in range(rng.randint(0, 8)))
        made = whole_characters(text)
        assert made == _made_whole(text), ascii(text)
        made.encode("utf-8")


def test_text_with_no_half_in_it_is_given_back_as_it_is():
    for value in itertools.chain(every_short_value(), longer_values(), FORMULAS, TEXT):
        assert whole_characters(value) is value


def _halved(value: str):
    """value with half of a character put in at each place, the high half
    and the low one, and with a pair that came apart at its end."""
    for i in range(len(value) + 1):
        for half in (HIGH, LOW):
            yield value[:i] + half + value[i:]
    yield value + HIGH + LOW


def test_every_writer_makes_a_half_whole_before_its_rule():
    """A provider's JSON read by Python's json module, and a statement's
    text read by pdfplumber, can hold half of a character, which UTF-8
    cannot write. Every writer gives for such a text exactly what it gives
    for the text made whole, nothing it gives holds a half, and the apps'
    own files read it back made whole."""
    checked = 0
    for value in itertools.chain(FORMULAS, TEXT, PLAIN_NEGATIVE, [p for p, _ in PIECES]):
        for text in _halved(value):
            made = _made_whole(text)
            assert csv_value(text) == csv_value(made), ascii(text)
            assert store_value(text) == store_value(made), ascii(text)
            assert xlsx_value(text) == xlsx_value(made), ascii(text)
            assert read_stored(store_value(text)) == made, ascii(text)
            for written in (csv_value(text), store_value(text)):
                written.encode("utf-8")
                assert not live_anywhere(written), ascii(text)
            xlsx_value(text).encode("utf-8")
            checked += 1
    assert checked > 1000


# A value holding half a character, and what one of the apps' own files
# holds for it, U+FFFD where the half was. Both tools write the same.
HALF_STORED = [
    ("SHOP " + HIGH + "NAME", "SHOP " + MARK + "NAME"),
    ("Great gift " + chr(0xD83C), "Great gift " + MARK),
    (HIGH + LOW + " Kettle", chr(0x1F600) + " Kettle"),
    (HIGH + "=1+1", MARK + "=1+1"),
    ("=1+1" + HIGH, "'=1+1" + MARK),
    ("x;" + HIGH + "=1+1", "x;" + MARK + "=1+1"),
    ("x;=1+1" + LOW, "x;'=1+1" + MARK),
    ("'" + HIGH + "=1+1", "'" + MARK + "=1+1"),
    ("=3+3 Gift " + chr(0xD83C) + ".pdf", "'=3+3 Gift " + MARK + ".pdf"),
]


@pytest.mark.parametrize("value, held", HALF_STORED)
def test_what_the_file_holds_for_half_a_character_and_what_reads_back(value, held):
    assert store_value(value) == held and csv_value(value) == held
    assert read_stored(held) == _made_whole(value)
    held.encode("utf-8")
    assert not live_anywhere(held)


# -- the census -------------------------------------------------------------------

# A tool in tools/ that writes a spreadsheet and need not go through the
# rule, with the reason. Anything else that writes one has to.
NOT_FOR_PEOPLE = {
    "make_sample.py": "writes the sample archive's own invented records, in the "
                      "format the apps write, for the exporters to read",
    "snapshot_traffic.py": "writes the repository's daily view and clone counts, "
                           "dates and numbers only",
}
# What each kind of spreadsheet has to be written through.
NEEDS = {"csv": {"csv_value"}, "xlsx": {"xlsx_value", "keep_text"}}


def _kinds(tree) -> set:
    """Which spreadsheets a tool writes, a CSV through the csv module's
    writers and an xlsx through openpyxl's Workbook."""
    kinds = set()
    for node in ast.walk(tree):
        if isinstance(node, ast.Attribute) and node.attr in ("writer", "DictWriter") \
                and isinstance(node.value, ast.Name) and node.value.id == "csv":
            kinds.add("csv")
        if isinstance(node, ast.ImportFrom) and (node.module or "").startswith("openpyxl") \
                and any(a.name == "Workbook" for a in node.names):
            kinds.add("xlsx")
    return kinds


def _imported_from_the_rule(tree) -> set:
    names = set()
    for node in ast.walk(tree):
        if isinstance(node, ast.ImportFrom) and node.module == "paperpull_core.sheets":
            names |= {a.name for a in node.names}
    return names


def _called(node) -> set:
    return {n.func.id for n in ast.walk(node)
            if isinstance(n, ast.Call) and isinstance(n.func, ast.Name)}


def _writing_functions(tree) -> list:
    """Each function that makes a CSV writer or an openpyxl Workbook, the
    functions it holds included, with what it writes."""
    return [(node, _kinds(node)) for node in ast.walk(tree)
            if isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef)) and _kinds(node)]


def writers() -> dict:
    out = {}
    for p in sorted(TOOLS.glob("*.py")):
        tree = ast.parse(p.read_text(encoding="utf-8"))
        if _kinds(tree):
            out[p.name] = tree
    return out


def test_the_census_finds_both_spreadsheet_tools_and_what_each_writes():
    found = {name: _kinds(tree) for name, tree in writers().items()}
    assert found["export_purchases.py"] == {"csv", "xlsx"}
    assert found["export_transactions.py"] == {"csv", "xlsx"}
    names = {name: sorted(f.name for f, _ in _writing_functions(tree))
             for name, tree in writers().items()}
    assert names["export_purchases.py"] == ["write_csv", "write_xlsx"]
    assert names["export_transactions.py"] == ["_write_csv", "_write_xlsx"]


@pytest.mark.parametrize("name", sorted(set(writers()) - set(NOT_FOR_PEOPLE)))
def test_every_function_writing_a_spreadsheet_for_a_person_goes_through_the_rule(name):
    tree = writers()[name]
    rule = _imported_from_the_rule(tree)
    functions = _writing_functions(tree)
    assert functions, "%s writes a spreadsheet outside any function" % name
    for func, kinds in functions:
        missing = set().union(*(NEEDS[k] for k in kinds)) - (rule & _called(func))
        assert not missing, "%s's %s writes a spreadsheet without paperpull_core.sheets' %s" % (
            name, func.name, ", ".join(sorted(missing)))


@pytest.mark.parametrize("tool", ["export_purchases.py", "export_transactions.py"])
def test_each_tool_takes_the_core_beside_it_with_no_core_installed(tool, tmp_path):
    """-S leaves out every installed package, so the only core either tool
    can import is the one in core/ beside tools/, which is where a checkout,
    the packaged builds and the server image all keep it. The tool runs as
    a script, as the transactions export runs, and is loaded by its path,
    as the panel loads the purchases export, which puts no folder on
    sys.path."""
    path = TOOLS / tool
    ran = subprocess.run([sys.executable, "-I", "-S", str(path), "--help"], cwd=tmp_path,
                         capture_output=True, text=True, timeout=120)
    assert ran.returncode == 0 and "usage" in ran.stdout, ran.stderr
    load = ("import importlib.util, sys\n"
            "spec = importlib.util.spec_from_file_location('tool', sys.argv[1])\n"
            "mod = importlib.util.module_from_spec(spec)\n"
            "spec.loader.exec_module(mod)\n"
            "print(sys.modules['paperpull_core.sheets'].__file__)\n")
    loaded = subprocess.run([sys.executable, "-I", "-S", "-c", load, str(path)], cwd=tmp_path,
                            capture_output=True, text=True, timeout=120)
    assert loaded.returncode == 0, loaded.stderr
    assert Path(loaded.stdout.strip()).resolve() == \
        (TOOLS.parent / "core" / "paperpull_core" / "sheets.py").resolve()


def test_a_tool_set_aside_still_writes_one_and_says_why():
    assert set(NOT_FOR_PEOPLE) <= set(writers())
    assert all(why.strip() for why in NOT_FOR_PEOPLE.values())


# -- the apps' own CSV files ---------------------------------------------------------

REPO = Path(__file__).resolve().parents[2]
# A minus, at most one of the money signs that are not letters, and a plain
# number, which Excel opened as a number or as text and never as a formula.
# Written out again here, so the model below is not the rule checking itself.
_EXCEL_PLAIN = re.compile(r"-[$£€]?(?:[0-9]{1,3}(?:,[0-9]{3})+|[0-9]+)(?:\.[0-9]+)?")
_BREAK = re.compile(r";|\r\n|\r|\n")


def _live(cell: str) -> bool:
    """Whether a spreadsheet would work a cell out, seen from its side.
    Excel drops a NUL from a CSV it opens, a spreadsheet may drop blank
    space in front, and then a sign starts a formula unless the rest is a
    plain negative number."""
    lead = cell.replace("\x00", "").lstrip()
    return (bool(lead) and unicodedata.normalize("NFKC", lead[0])[:1] in "=+-@"
            and not _EXCEL_PLAIN.fullmatch(lead))


def _split_at_semicolons(piece: str) -> list:
    """What a reader splitting a CSV at ; makes of a piece after a ; or a
    line break, with and without blank space in front. The csv module
    writes each " of a quoted value doubled."""
    return _as_a_semicolon_reader_reads(piece.replace('"', '""'))


def _as_a_semicolon_reader_reads(token: str) -> list:
    """The cells a reader splitting at ; may make of a token as the file
    holds it, with and without blank space in front. A cell that starts
    with " is a quoted part, "" a quote inside it, with the rest after it."""
    token = token.replace("\x00", "")
    cells = []
    for t in (token, token.lstrip()):
        if not t.startswith('"'):
            cells.append(t)
            continue
        i, quoted = 1, []
        while i < len(t):
            if t[i] == '"':
                if t[i + 1:i + 2] == '"':
                    quoted.append('"')
                    i += 2
                    continue
                i += 1
                break
            quoted.append(t[i])
            i += 1
        cells.append("".join(quoted) + t[i:])
    return cells


def live_anywhere(text: str) -> bool:
    """The whole value as Excel opens a comma CSV, or any piece of it a
    reader splitting at ; or a line break makes a cell of its own. A model
    of the spreadsheet, not of the rule."""
    return _live(text) or any(_live(c) for p in _BREAK.split(text)[1:]
                              for c in _split_at_semicolons(p))


def live_in_a_row(cells) -> bool:
    """A row as the csv module writes it, read whole by a spreadsheet in a
    country whose list separator is ; and whose decimal sign is a comma. It
    splits the line at ; and line breaks, so the last piece of one value
    runs on through the comma that ends it into the next value, and -1
    beside 5+1 is read as -1,5+1."""
    buf = io.StringIO()
    csv.writer(buf).writerow(cells)
    return any(_live(c) for token in _BREAK.split(buf.getvalue())
               for c in _as_a_semicolon_reader_reads(token))


# Values of other people's choosing to set side by side in a row. The first
# cell of every file the apps and the tools write is the person's own name
# or the provider's, so a row here starts with one.
NEIGHBORS = ["-1", "q;-1", "q; -1", "q;-$12.34", "-1;q", "5+1", "5+1;q", "5", "2026-09-12",
             "-$12.34", "Refund; -12.34", "x;", ";", "=1+1", "q;=1+1", 'q;"-1', "q;'-1",
             "q;\x00-1", "' -1", "@x", "Widget", ""]


@pytest.mark.parametrize("write", [store_value, csv_value], ids=["store_value", "csv_value"])
def test_no_value_runs_on_into_the_next_as_a_formula(write):
    assert live_in_a_row(["Dana Example", "q;-1", "5+1;q"]), "the model misses what it is for"
    for left in NEIGHBORS:
        for right in NEIGHBORS:
            row = ["Dana Example", write(left), write(right)]
            assert not live_in_a_row(row), (left, right)


# Measured in Excel 365 on 2026-10-16, each written as it is, the way the
# rule of the first commit of this change let it through. A NUL in front
# opened as a working formula. A piece starting with " after a ; opened as
# one in Excel's own import split at ;.
EXCEL_RAN = ["\x00=1+1", "\x00-1+1", "\x00@SUM(1,1)", "\x00+1+1",
             'x;"=1+1;y', 'x;"@SUM(1,1);y', 'x;"-1+1;y', 'x\n"=1+1;y', 'x\r\n"=1+1;y']


@pytest.mark.parametrize("text", EXCEL_RAN)
def test_the_model_knows_what_excel_ran_and_the_rule_keeps_it_text(text):
    assert live_anywhere(text)
    assert not live_anywhere(store_value(text)) and not live_anywhere(csv_value(text))
    assert read_stored(store_value(text)) == text


# What can sit in a value, each kind of thing the rule or a spreadsheet
# looks at.
ALPHABET = ["'", '"', "\x00", "=", "+", "-", "@", " ", "\t", "\r", "\n", ";", "1", "a",
            "＝", "$"]


def every_short_value():
    for n in range(5):
        for chars in itertools.product(ALPHABET, repeat=n):
            yield "".join(chars)


def longer_values(count=20000):
    pieces = ALPHABET + ["\r\n", "''", '""', "-$12.34", "-12.34", "=HYPERLINK(\"x\")", ",",
                         "é", "£", "kr", "SUM(1,1)", "Statement"]
    rnd = random.Random(20261016)
    for _ in range(count):
        yield "".join(rnd.choice(pieces) for _ in range(rnd.randint(0, 16)))


def test_every_value_reads_back_exactly_as_it_was_written():
    checked = 0
    for value in itertools.chain(every_short_value(), longer_values()):
        assert read_stored(store_value(value)) == value, repr(value)
        checked += 1
    assert checked > 50000


def test_nothing_written_is_ever_worked_out_by_a_spreadsheet():
    for value in itertools.chain(every_short_value(), longer_values()):
        assert not live_anywhere(store_value(value)), repr(value)
        assert not live_anywhere(csv_value(value)), repr(value)


def test_a_value_with_no_quote_in_it_is_written_as_the_exporters_write_it():
    """One rule. A value only differs from the exporters' CSV where a quote
    of either kind starts it, which the exporters need not tell apart since
    nothing reads their file back."""
    for value in itertools.chain(every_short_value(), longer_values()):
        if "'" not in value and '"' not in value:
            assert store_value(value) == csv_value(value), repr(value)


# A value, and what the file holds for it. These are the files' format,
# which may grow and never shrink (sheets, THE APPS' OWN CSV FILES), since a
# file written by one version is read by the next.
STORED = [
    ("=1+1", "'=1+1"),
    ('=HYPERLINK("https://example.invalid/?"&A1,"open")',
     '\'=HYPERLINK("https://example.invalid/?"&A1,"open")'),
    ("@SUM(1,1)", "'@SUM(1,1)"),
    ("-1+1", "'-1+1"),
    ("=3+3; @SUM(1,1) Statement.pdf", "'=3+3;' @SUM(1,1) Statement.pdf"),
    # A real name that starts with a quote and then = is told apart from one
    # given its quote here by one more quote.
    ("'=1+1", "''=1+1"),
    ("''=1+1", "'''=1+1"),
    ("' =1+1", "'' =1+1"),
    ("Widget;'=1+1", "Widget;''=1+1"),
    ("\n=1+1", "'\n'=1+1"),
    (" =1+1", "' =1+1"),
    ("＝1+1", "'＝1+1"),
    ("-$A$1", "'-$A$1"),
    ("-kr12.34", "'-kr12.34"),
    # A NUL Excel drops, and a piece starting with a " that a reader
    # splitting at ; takes for an empty quoted part.
    ("\x00=1+1", "'\x00=1+1"),
    ("x;\x00=1+1", "x;'\x00=1+1"),
    ('x;"=1+1;y', 'x;\'"=1+1;y'),
    ('x\r\n"@SUM(1,1)', 'x\r\n\'"@SUM(1,1)'),
    ('x; "=1+1', 'x;\' "=1+1'),
    ('"=1+1', '\'"=1+1'),
    ('x;"\'=1+1', 'x;\'"\'=1+1'),
    # Written as they are.
    ("'Tis the season ornament", "'Tis the season ornament"),
    ("O'Brien", "O'Brien"),
    ("'", "'"),
    ("''", "''"),
    ("Rock 'n' roll; 'live'", "Rock 'n' roll; 'live'"),
    ('12" deluxe; "special" edition', '12" deluxe; "special" edition'),
    ("-$12.34", "-$12.34"),
    ("-$1,234.56", "-$1,234.56"),
    ("-€12.34", "-€12.34"),
    ("-12.34", "-12.34"),
    ("Refund; -$12.34", "Refund;' -$12.34"),
    ("C:\\Users\\Dana\\Shop;=1+1\\a.pdf", "C:\\Users\\Dana\\Shop;'=1+1\\a.pdf"),
    ("", ""),
]


@pytest.mark.parametrize("value, held", STORED)
def test_what_the_file_holds_and_what_reads_back(value, held):
    assert store_value(value) == held
    assert read_stored(held) == value
    assert not live_anywhere(held)


# What a file written before the rule could hold, which reads as it did.
OLD = ["=1+1", "@SUM(1,1)", "-1+1", "+1+1", " =1+1", "Widget;=1+1", "\n=1+1",
       "O'Brien", "'Tis the season ornament", "Rock 'n' roll; 'live'", "'", "''",
       "'plain words", "a;'b", "-$12.34", "-12.34", "'-12.34", "'-$12.34",
       "2026-09-12", "https://example.invalid/r?id=1", "C:\\Users\\Dana\\a.pdf"]


@pytest.mark.parametrize("value", OLD)
def test_a_file_written_before_the_rule_reads_as_it_did(value):
    assert read_stored(value) == value


def test_the_one_old_shape_that_reads_differently_is_a_quote_before_a_formula():
    """A value an older version wrote starting with a quote and then text
    like a formula reads without that quote, since that is what a quote
    put in by the rule looks like. The real installs held no such value on
    2026-10-16, and no value starting with a quote at all."""
    assert read_stored("'=1+1") == "=1+1"
    assert read_stored("Widget;'=1+1") == "Widget;=1+1"
    assert read_stored("'\"=1+1") == '"=1+1'
    for value in itertools.chain(every_short_value(), longer_values()):
        if not any(p.startswith("'") for p in [value] + _BREAK.split(value)[1:]):
            assert read_stored(value) == value, repr(value)


def test_none_is_an_empty_cell_and_a_number_is_written_as_its_text():
    assert store_value(None) == ""
    assert (store_value(-5), store_value(12.5), store_value(0), store_value(True)) == \
        ("-5", "12.5", "0", "True")
    # Text a number gives that a spreadsheet could take for a formula gets
    # its quote too, and reads back as that text.
    assert store_value(-1e-05) == "'-1e-05" and read_stored("'-1e-05") == "-1e-05"


def test_what_is_not_text_reads_back_as_it_is():
    """csv.DictReader hands None for a short row's missing values and a
    list under None for a long row's extra ones."""
    assert read_stored(None) is None
    row = read_stored_row({"Item Name": "'=1+1", "Notes": None, None: ["'=1+1"]})
    assert row == {"Item Name": "=1+1", "Notes": None, None: ["'=1+1"]}


# -- the census of every CSV reader and writer ------------------------------------------

ROOTS = ("apps", "core", "gui", "server", "tools")
MADE_FOLDERS = {"build", "dist", "site-packages", "node_modules"}


def sources():
    """Every Python file of the repository that is not a test, a copy a
    build made or a linked environment."""
    out = []
    for root in ROOTS:
        for folder, dirs, files in os.walk(REPO / root):
            dirs[:] = sorted(d for d in dirs if d != "tests" and d not in MADE_FOLDERS
                             and not d.startswith((".", "__")) and not d.endswith(".egg-info"))
            out.extend(Path(folder) / f for f in files
                       if f.endswith(".py") and not f.startswith("test_") and f != "conftest.py")
    return sorted(out)


@functools.lru_cache(maxsize=None)
def _parsed(repo: Path) -> tuple:
    """Each source file and its tree, read once for every test here."""
    return tuple((p, ast.parse(p.read_text(encoding="utf-8-sig"))) for p in sources())


def _kind(n):
    """'reads' or 'writes' for a csv reader or writer made at this node, a
    pandas-style read_csv and to_csv counted too, else None."""
    if not isinstance(n, ast.Attribute):
        return None
    if n.attr == "read_csv":
        return "reads"
    if n.attr == "to_csv":
        return "writes"
    if isinstance(n.value, ast.Name) and n.value.id == "csv":
        if n.attr in ("reader", "DictReader"):
            return "reads"
        if n.attr in ("writer", "DictWriter"):
            return "writes"
    return None


def _makes(node) -> set:
    """Whether a function makes a csv reader or a csv writer."""
    return {k for k in map(_kind, ast.walk(node)) if k}


def _calls(node) -> set:
    names = set()
    for n in ast.walk(node):
        if isinstance(n, ast.Call):
            f = n.func
            names.add(f.id if isinstance(f, ast.Name) else f.attr if isinstance(f, ast.Attribute)
                      else "")
    return names


def csv_functions() -> dict:
    """'<path>::<function>' to what it does with a csv, for every function
    in the repository that makes a csv reader or writer."""
    out = {}
    for path, tree in _parsed(REPO):
        for node in ast.walk(tree):
            if isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef)) and _makes(node):
                name = "%s::%s" % (path.relative_to(REPO).as_posix(), node.name)
                out[name] = (_makes(node), _calls(node))
    return out


def csv_out_of_sight() -> list:
    """Every way of reaching csv the census above would not see. A reader or
    writer made outside any function, and csv imported under another name
    or a name taken from it. Code splitting the files' text itself is not
    seen by any of this, and on 2026-10-16 a search of the repository for
    the three files' names and paths found none."""
    out = []
    for path, tree in _parsed(REPO):
        rel = path.relative_to(REPO).as_posix()
        inside = {id(n) for f in ast.walk(tree)
                  if isinstance(f, (ast.FunctionDef, ast.AsyncFunctionDef)) for n in ast.walk(f)}
        for n in ast.walk(tree):
            if _kind(n) and id(n) not in inside:
                out.append("%s:%d makes a csv %s outside a function" % (rel, n.lineno, _kind(n)))
            if isinstance(n, ast.Import) and any(a.name == "csv" and a.asname for a in n.names):
                out.append("%s:%d imports csv under another name" % (rel, n.lineno))
            if isinstance(n, ast.ImportFrom) and n.module == "csv":
                out.append("%s:%d takes a name from csv" % (rel, n.lineno))
    return out


def test_no_csv_reader_or_writer_is_hidden_from_the_census():
    assert csv_out_of_sight() == []


# How each kind of use has to go through the rule. A writer of an app's own
# file has to use store_value, since csv_value's quote does not always read
# back, read_stored("'=1+1") of a real '=1+1 among them. The two tools write
# files nobody reads back, through csv_value.
THROUGH = {"reads": {"read_stored", "read_stored_row"}, "writes": {"store_value"}}
EXPORTS = {"tools/export_purchases.py::write_csv", "tools/export_transactions.py::_write_csv"}
# A function that makes a csv reader or writer for something other than an
# app's own record, with the reason. Anything else has to go through the rule.
NOT_A_RECORD = {
    "core/paperpull_core/storage.py::_ensure_header":
        "writes the header row alone, the app's own column names",
    "tools/make_sample.py::statements":
        "writes the sample archive's invented records, each value checked below "
        "to be held as an app would hold it",
    "tools/make_sample.py::receipts":
        "writes the sample archive's invented records, each value checked below "
        "to be held as an app would hold it",
    "tools/snapshot_traffic.py::merge":
        "reads and writes the repository's daily view and clone counts, dates and numbers",
}
# What the census has to find, so a change that hides one fails here.
MUST_FIND = {
    "core/paperpull_core/storage.py::append_rows": {"writes"},
    "core/paperpull_core/storage.py::read_all": {"reads"},
    "core/paperpull_core/storage.py::rewrite": {"writes"},
    "tools/export_purchases.py::load_purchases": {"reads"},
    "tools/export_transactions.py::_index_rows": {"reads"},
    "tools/export_purchases.py::write_csv": {"writes"},
    "tools/export_transactions.py::_write_csv": {"writes"},
}


def test_the_census_finds_every_reader_and_writer_of_an_apps_csv():
    found = csv_functions()
    for name, kinds in MUST_FIND.items():
        assert name in found and found[name][0] == kinds, name
    assert set(NOT_A_RECORD) <= set(found), sorted(set(NOT_A_RECORD) - set(found))
    assert all(why.strip() for why in NOT_A_RECORD.values())


@pytest.mark.parametrize("name", sorted(set(csv_functions()) - set(NOT_A_RECORD) - EXPORTS))
def test_every_csv_reader_and_writer_goes_through_the_rule(name):
    kinds, calls = csv_functions()[name]
    for kind in kinds:
        assert THROUGH[kind] & calls, "%s %s a CSV without paperpull_core.sheets' %s" % (
            name, kind, " or ".join(sorted(THROUGH[kind])))


@pytest.mark.parametrize("name", sorted(EXPORTS))
def test_each_tools_own_csv_goes_through_csv_value(name):
    kinds, calls = csv_functions()[name]
    assert kinds == {"writes"} and "csv_value" in calls, name


def _load_tool(name: str):
    spec = importlib.util.spec_from_file_location(name, TOOLS / ("%s.py" % name))
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


def test_the_sample_archive_holds_each_value_as_an_app_would(tmp_path):
    """The sample writes its records with csv itself. Each value it writes
    has to be one the rule leaves as it is, so every reader reads the
    sample as written."""
    _load_tool("make_sample").build(tmp_path)
    files = [p for p in tmp_path.rglob("*.csv")
             if p.name.endswith((" Order History.csv", " Receipt Index.csv", " Document Index.csv"))]
    assert len(files) >= 3
    values = 0
    for path in files:
        with open(path, encoding="utf-8-sig", newline="") as fh:
            for row in csv.DictReader(fh):
                for value in row.values():
                    assert store_value(value) == value, (path.name, value)
                    values += 1
    assert values > 100
